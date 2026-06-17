#!/usr/bin/env python3
"""
preencher_tabela.py
Lê resultados.txt gerado por rodar_crivo.sh e preenche Crivo_organizado.xlsx.

Uso:
    python preencher_tabela.py [resultados.txt] [Crivo.xlsx] [Crivo_preenchido.xlsx]

Valores de N suportados: 10^8 até 10^12 (5 colunas).
"""

import sys
import os
from openpyxl import load_workbook

ARQ_RESULTADOS = sys.argv[1] if len(sys.argv) > 1 else "resultados.txt"
ARQ_ENTRADA    = sys.argv[2] if len(sys.argv) > 2 else "Crivo_organizado.xlsx"
ARQ_SAIDA      = sys.argv[3] if len(sys.argv) > 3 else "Crivo_preenchido.xlsx"

# Bloco esquerdo: Tempo e Memória (cols C–G)
COL_TEMPO = {
    100_000_000:         "C",
    1_000_000_000:       "D",
    10_000_000_000:      "E",
    100_000_000_000:     "F",
    1_000_000_000_000:   "G",
}

# Bloco direito: Primos e Marcações (cols L–P)
COL_DIR = {
    100_000_000:         "L",
    1_000_000_000:       "M",
    10_000_000_000:      "N",
    100_000_000_000:     "O",
    1_000_000_000_000:   "P",
}

# Tempo: BOOL linhas 3-12, BITSET linhas 15-24
LINHA_BASE_TEMPO_BOOL   = 3
LINHA_BASE_TEMPO_BITSET = 15

# Memória: BOOL linhas 29-38, BITSET linhas 41-50
LINHA_BASE_MEM_BOOL   = 29
LINHA_BASE_MEM_BITSET = 41

# Primos: BOOL linhas 3-12, BITSET linhas 15-24 (bloco direito)
LINHA_BASE_PRIMOS_BOOL   = 3
LINHA_BASE_PRIMOS_BITSET = 15

# Marcações: linha 55 (BOOL) e 56 (BITSET)
LINHA_MARC_BOOL   = 55
LINHA_MARC_BITSET = 56

def ler_resultados(caminho):
    dados = {}
    if not os.path.exists(caminho):
        print(f"ERRO: arquivo '{caminho}' não encontrado.")
        sys.exit(1)
    with open(caminho) as f:
        for linha in f:
            linha = linha.strip()
            if not linha or linha.startswith("#"):
                continue
            partes = linha.split()
            if len(partes) < 7:
                continue
            prog, N, rodada, tempo, primos, marcacoes, mem_kb = partes[:7]
            N = int(N)
            rodada = int(rodada)
            if N not in COL_TEMPO:
                continue  # ignora N fora do escopo (ex: 10^13 ou 10^14)
            prog_base = "BOOL" if "BITSET" not in prog.upper() else "BITSET"
            dados[(prog_base, N, rodada)] = {
                "tempo":     float(tempo),
                "primos":    int(primos),
                "marcacoes": int(marcacoes),
                "mem_mb":    float(mem_kb) / 1024.0,
            }
    return dados

def preencher(dados, arq_in, arq_out):
    wb = load_workbook(arq_in)
    ws = wb.active

    for N in sorted(COL_TEMPO.keys()):
        for rodada in range(1, 11):
            for prog in ("BOOL", "BITSET"):
                chave = (prog, N, rodada)
                if chave not in dados:
                    continue
                d = dados[chave]
                off = rodada - 1

                # Tempo
                col_t = COL_TEMPO[N]
                base_t = LINHA_BASE_TEMPO_BOOL if prog == "BOOL" else LINHA_BASE_TEMPO_BITSET
                ws[f"{col_t}{base_t + off}"] = round(d["tempo"], 6)

                # Memória
                base_m = LINHA_BASE_MEM_BOOL if prog == "BOOL" else LINHA_BASE_MEM_BITSET
                ws[f"{col_t}{base_m + off}"] = round(d["mem_mb"], 3)

                # Primos
                col_d = COL_DIR[N]
                base_p = LINHA_BASE_PRIMOS_BOOL if prog == "BOOL" else LINHA_BASE_PRIMOS_BITSET
                ws[f"{col_d}{base_p + off}"] = d["primos"]

            # Marcações (valor da última rodada processada para este N)
            for prog, linha_marc in (("BOOL", LINHA_MARC_BOOL),
                                     ("BITSET", LINHA_MARC_BITSET)):
                chave = (prog, N, rodada)
                if chave in dados:
                    ws[f"{COL_DIR[N]}{linha_marc}"] = dados[chave]["marcacoes"]

    wb.save(arq_out)
    print(f"Planilha salva em: {arq_out}")

if __name__ == "__main__":
    print(f"Lendo: {ARQ_RESULTADOS}")
    dados = ler_resultados(ARQ_RESULTADOS)
    print(f"Registros lidos: {len(dados)}")
    preencher(dados, ARQ_ENTRADA, ARQ_SAIDA)
